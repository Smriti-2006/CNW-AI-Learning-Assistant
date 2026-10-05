import streamlit as st
import json
import re
import pandas as pd
import joblib
from textwrap import dedent

from google import genai

# Optional voice input
try:
    from streamlit_mic_recorder import speech_to_text
    VOICE_AVAILABLE = True
except Exception:
    speech_to_text = None
    VOICE_AVAILABLE = False

# Optional PowerPoint support
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    PPT_AVAILABLE = True
except Exception:
    PPT_AVAILABLE = False


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="CNW AI Learning Assistant",
    page_icon="📡",
    layout="wide",
    initial_sidebar_state="expanded"
)


# ============================================================
# GEMINI CLIENT
# ============================================================

try:
    client = genai.Client(
        api_key=st.secrets["GEMINI_API_KEY"]
    )
except Exception:
    client = None


MODEL_NAME = "gemini-3.5-flash-lite"
qos_model = joblib.load("random_forest_qos_model_compressed.pkl")
preprocessor = joblib.load("qos_preprocessor_compatible.pkl")
scaler = joblib.load("qos_scaler.pkl")

# ============================================================
# SESSION STATE
# ============================================================

defaults = {
    "logged_in": False,
    "user_name": "",
    "selected_page": "🏠 Home",

    "history": [],

    # Study
    "study_question": "",
    "study_answer": "",
    "study_pending_question": "",
    "study_chat": [],

    # Exam
    "exam_question": "",
    "exam_answer": "",
    "exam_mark": "5 Marks",
    "exam_pending_question": "",
    "exam_chat": [],

    # Viva
    "viva_started": False,
    "viva_questions": [],
    "viva_index": 0,
    "viva_score": 0,
    "viva_answered": False,
    "viva_feedback": "",

    # Presentation
    "presentation_data": None,

    # Mind map
    "mindmap_data": None,

    # Flashcards
    "flashcards": [],
    "flashcard_index": 0,

    # Quiz
    "quiz_started": False,
    "quiz_questions": [],
    "quiz_index": 0,
    "quiz_score": 0,
    "quiz_answered": False,

    # Predictor
    "predictor_prediction": None
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def render_html(html):
    st.html(dedent(html))


def clean_json_text(text):
    """
    Clean Gemini response before JSON parsing.
    """

    if not text:
        return ""

    text = str(text).strip()

    text = text.replace("```json", "")
    text = text.replace("```JSON", "")
    text = text.replace("```", "")

    start = text.find("[")
    end = text.rfind("]")

    if start != -1 and end != -1:
        return text[start:end + 1]

    start = text.find("{")
    end = text.rfind("}")

    if start != -1 and end != -1:
        return text[start:end + 1]

    return text


def clean_ai_text(text):
    """
    Remove unwanted Markdown symbols from AI answers.
    """

    if not text:
        return ""

    text = str(text)

    # Remove Markdown headings
    text = re.sub(
        r"^\s*#{1,6}\s*",
        "",
        text,
        flags=re.MULTILINE
    )

    # Remove bold / italic markers
    text = text.replace("***", "")
    text = text.replace("**", "")
    text = text.replace("__", "")
    text = text.replace("~~", "")

    # Remove code fences
    text = text.replace("```python", "")
    text = text.replace("```json", "")
    text = text.replace("```JSON", "")
    text = text.replace("```", "")
    text = text.replace("`", "")

    # Remove horizontal lines
    text = re.sub(
        r"^\s*-{3,}\s*$",
        "",
        text,
        flags=re.MULTILINE
    )

    text = re.sub(
        r"^\s*_{3,}\s*$",
        "",
        text,
        flags=re.MULTILINE
    )

    # Remove excessive blank lines
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text
    )

    return text.strip()


def format_ai_text(text):
    """
    Safely convert AI text into HTML.
    """

    text = clean_ai_text(text)

    text = (
        text
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

    text = text.replace("\n", "<br>")

    return text


def ask_gemini(prompt):

    if client is None:
        return (
            "Gemini connection is not available. "
            "Please check your GEMINI_API_KEY in "
            ".streamlit/secrets.toml."
        )

    try:

        response = client.interactions.create(
            model=MODEL_NAME,
            input=prompt
        )

        return response.output_text

    except Exception as e:

        return (
            f"Unable to generate the answer right now.\n\n"
            f"Error: {str(e)}"
        )


def voice_input(key, label="🎤"):
    """Convert spoken English into text using the browser microphone."""
    if not VOICE_AVAILABLE:
        st.warning("Voice input needs streamlit-mic-recorder. Run: pip install streamlit-mic-recorder")
        return None
    try:
        return speech_to_text(
            language="en",
            start_prompt=label,
            stop_prompt="⏹️",
            just_once=True,
            use_container_width=True,
            key=key
        )
    except Exception as e:
        st.error(f"Voice input could not start: {e}")
        return None


def chat_input_with_voice(text_key, voice_key, placeholder, button_key, button_text="Send"):
    """ChatGPT-style input row with a microphone beside the text box."""
    left, right = st.columns([0.91, 0.09], gap="small", vertical_alignment="bottom")

    # Call voice widget before creating the text widget so session state can be updated safely.
    with right:
        voice_text = voice_input(voice_key, "🎤")

    if voice_text:
        st.session_state[text_key] = voice_text

    with left:
        text_value = st.text_input(
            "",
            placeholder=placeholder,
            key=text_key,
            label_visibility="collapsed"
        )

    submitted = st.button(button_text, use_container_width=True, key=button_key)
    return text_value, submitted


def add_history(mode, topic):

    st.session_state.history.insert(
        0,
        {
            "mode": mode,
            "topic": topic
        }
    )

    # Keep only latest 30 entries
    st.session_state.history = (
        st.session_state.history[:30]
    )


# ============================================================
# GLOBAL CSS
# ============================================================

st.markdown(
    """
    <style>

    /* =====================================================
       GLOBAL
       ===================================================== */

    .stApp {
        background: #F5F9FF;
        color: #172033;
    }

    html, body, [class*="css"] {
        font-size: 17px;
    }

    p {
        font-size: 17px;
        line-height: 1.65;
    }

    .main .block-container {
        padding-top: 1.5rem;
        padding-bottom: 3rem;
        max-width: 1500px;
    }


    /* =====================================================
       SIDEBAR
       ===================================================== */

    [data-testid="stSidebar"] {
        background: #FFFFFF;
        border-right: 1px solid #DCE6F2;
    }

    [data-testid="stSidebar"] > div:first-child {
        padding-top: 1.5rem;
    }

    [data-testid="stSidebar"] .stButton > button {
        background: transparent !important;
        color: #172033 !important;
        border: none !important;
        text-align: left !important;
        box-shadow: none !important;
        font-size: 16px !important;
        font-weight: 500 !important;
        border-radius: 12px !important;
        padding: 12px 15px !important;
    }

    [data-testid="stSidebar"] .stButton > button:hover {
        background: #EFF6FF !important;
        color: #2563EB !important;
    }

    [data-testid="stSidebar"] .stButton > button:focus {
        background: #DBEAFE !important;
        color: #1D4ED8 !important;
    }


    /* =====================================================
       TITLES
       ===================================================== */

    .page-title {
        font-size: 38px;
        font-weight: 850;
        color: #172033;
        margin-bottom: 8px;
        line-height: 1.15;
    }

    .page-subtitle {
        font-size: 20px;
        color: #64748B;
        margin-bottom: 28px;
        line-height: 1.5;
    }


    /* =====================================================
       HOME MODE CARDS
       ===================================================== */

    .mode-card {
        background: #FFFFFF;
        border: 1px solid #DCE6F2;
        border-radius: 22px;
        padding: 28px;
        min-height: 190px;
        transition:
            transform 0.25s ease,
            box-shadow 0.25s ease,
            border-color 0.25s ease;
        box-shadow:
            0 5px 16px rgba(23, 32, 51, 0.06);
        margin-bottom: 18px;
    }

    .mode-card:hover {
        transform: translateY(-8px) scale(1.02);
        border-color: #2563EB;
        box-shadow:
            0 18px 35px rgba(37, 99, 235, 0.18);
    }

    .mode-icon {
        font-size: 42px;
        margin-bottom: 12px;
    }

    .mode-title {
        font-size: 23px;
        font-weight: 800;
        color: #172033;
        margin-bottom: 9px;
    }

    .mode-description {
        color: #64748B;
        font-size: 16px;
        line-height: 1.55;
    }


    /* =====================================================
       ANSWER CARD
       ===================================================== */

    .answer-card {
        background: #FFFFFF;
        border: 1px solid #DCE6F2;
        border-left: 6px solid #2563EB;
        border-radius: 20px;
        padding: 28px;
        margin-top: 20px;
        margin-bottom: 20px;
        box-shadow:
            0 8px 24px rgba(23, 32, 51, 0.07);
    }

    .answer-title {
        font-size: 24px;
        font-weight: 800;
        color: #2563EB;
        margin-bottom: 18px;
    }

    .answer-content {
        color: #172033;
        font-size: 18px;
        line-height: 1.8;
    }


    /* =====================================================
       INFO CARD
       ===================================================== */

    .info-card {
        background: #FFFFFF;
        border: 1px solid #DCE6F2;
        border-radius: 20px;
        padding: 28px;
        margin: 15px 0;
        box-shadow:
            0 6px 20px rgba(23, 32, 51, 0.06);
    }


    /* =====================================================
       WELCOME
       ===================================================== */

    .welcome-shell {
        min-height: calc(100vh - 80px);
        display: flex;
        align-items: center;
        justify-content: center;
        padding: 20px;
        box-sizing: border-box;
    }

    .welcome-card {
        width: 100%;
        max-width: 1050px;
        background: #FFFFFF;
        border-radius: 30px;
        border: 1px solid #DCE6F2;
        overflow: hidden;
        box-shadow:
            0 25px 60px rgba(37, 99, 235, 0.13);
    }

    .welcome-left {
        min-height: 500px;
        background:
            linear-gradient(
                135deg,
                #2563EB,
                #06B6D4
            );
        color: white;
        padding: 48px;
        display: flex;
        flex-direction: column;
        justify-content: center;
    }

    .welcome-icon {
        font-size: 58px;
        margin-bottom: 12px;
    }

    .welcome-brand {
        font-size: 25px;
        font-weight: 800;
        margin-bottom: 5px;
    }

    .welcome-heading {
        font-size: 40px;
        font-weight: 850;
        margin-bottom: 15px;
    }

    .welcome-description {
        font-size: 18px;
        line-height: 1.7;
        max-width: 520px;
        opacity: 0.95;
    }

    .feature-pill-row {
        display: flex;
        flex-wrap: wrap;
        gap: 10px;
        margin-top: 25px;
    }

    .feature-pill {
        background: rgba(255,255,255,0.18);
        border: 1px solid rgba(255,255,255,0.3);
        padding: 9px 15px;
        border-radius: 999px;
        font-size: 14px;
    }

    .welcome-right {
        min-height: 500px;
        background: #FFFFFF;
        padding: 48px;
        display: flex;
        flex-direction: column;
        justify-content: center;
    }

    .signin-title {
        font-size: 34px;
        font-weight: 850;
        color: #172033;
        margin-bottom: 10px;
    }

    .signin-description {
        font-size: 17px;
        color: #64748B;
        margin-bottom: 25px;
    }


    /* =====================================================
       BUTTONS
       ===================================================== */

    .stButton > button {
        min-height: 48px !important;
        border-radius: 12px !important;
        font-size: 17px !important;
        font-weight: 650 !important;
        border: 1px solid #DCE6F2 !important;
    }

    .stButton > button:hover {
        border-color: #2563EB !important;
        color: #2563EB !important;
    }


    /* =====================================================
       INPUTS
       ===================================================== */

    input, textarea {
        font-size: 17px !important;
    }

    label {
        font-size: 16px !important;
    }



    /* Chat-style input + microphone */
    div[data-testid="stHorizontalBlock"] {
        align-items: end;
    }

    div[data-testid="stHorizontalBlock"] input {
        border-radius: 14px !important;
        min-height: 48px !important;
        border: 1px solid #DCE6F2 !important;
    }

    /* =====================================================
       FLASHCARD
       ===================================================== */

    .flashcard {
        border-radius: 28px;
        padding: 35px;
        min-height: 330px;
        background:
            linear-gradient(
                135deg,
                #EFF6FF,
                #ECFEFF
            );
        border: 2px solid #BFDBFE;
        box-shadow:
            0 15px 35px rgba(37,99,235,0.10);
    }

    .flashcard-number {
        font-size: 15px;
        font-weight: 800;
        color: #2563EB;
        margin-bottom: 18px;
    }

    .flashcard-title {
        font-size: 28px;
        font-weight: 850;
        color: #172033;
        margin-bottom: 18px;
    }

    .flashcard-content {
        font-size: 18px;
        line-height: 1.75;
        color: #334155;
    }


    /* =====================================================
       RESPONSIVE
       ===================================================== */

    @media (max-width: 900px) {

        .page-title {
            font-size: 31px;
        }

        .page-subtitle {
            font-size: 17px;
        }

        .welcome-left,
        .welcome-right {
            min-height: auto;
            padding: 32px;
        }

        .welcome-heading {
            font-size: 32px;
        }

    }

    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# LOGIN PAGE
# ============================================================

if not st.session_state.logged_in:

    # ========================================================
    # LOGIN PAGE STYLE
    # ========================================================

    st.markdown(
        """
        <style>

        /* Hide sidebar */
        [data-testid="stSidebar"] {
            display: none !important;
        }

        /* Hide Streamlit header */
        [data-testid="stHeader"] {
            display: none !important;
        }

        /* Navy background */
        .stApp {
            background:
                linear-gradient(
                    135deg,
                    #123B7A 0%,
                    #075985 55%,
                    #087F9F 100%
                ) !important;
        }

        /* Main area */
        .main .block-container {
            max-width: 1320px !important;
            padding-top: 65px !important;
            padding-bottom: 65px !important;
            padding-left: 0 !important;
            padding-right: 0 !important;
        }

        /* Remove gap between cards */
        [data-testid="stHorizontalBlock"] {
            gap: 0 !important;
            align-items: stretch !important;
        }

        /* Both Streamlit columns */
        [data-testid="stColumn"] {
            padding: 0 !important;
            margin: 0 !important;
            min-height: 650px !important;
        }

        /* ====================================================
           LEFT BLUE CARD
           ==================================================== */

        [data-testid="stColumn"]:first-child {
            background:
                linear-gradient(
                    145deg,
                    #2563EB 0%,
                    #1688C4 50%,
                    #06B6D4 100%
                ) !important;

            border-radius: 30px 0 0 30px !important;

            box-shadow:
                0 25px 60px rgba(0,0,0,0.28) !important;

            overflow: hidden !important;

            position: relative !important;
        }

        /* Decorative circle */
        [data-testid="stColumn"]:first-child::after {
            content: "";

            position: absolute;

            width: 300px;
            height: 300px;

            border-radius: 50%;

            background:
                rgba(255,255,255,0.09);

            top: -110px;
            right: -80px;

            pointer-events: none;
        }

        /* ====================================================
           RIGHT WHITE CARD
           ==================================================== */

        [data-testid="stColumn"]:last-child {
            background: #FFFFFF !important;

            border-radius: 0 30px 30px 0 !important;

            box-shadow:
                0 25px 60px rgba(0,0,0,0.28) !important;

            overflow: hidden !important;

            position: relative !important;
        }

        /* ====================================================
           LEFT CONTENT
           ==================================================== */

        .login-left-content {
            position: relative;
            z-index: 5;

            padding: 55px;
            color: white;
        }

        .login-logo {
            width: 82px;
            height: 82px;

            border-radius: 24px;

            display: flex;
            align-items: center;
            justify-content: center;

            background:
                rgba(255,255,255,0.16);

            border:
                1px solid
                rgba(255,255,255,0.32);

            font-size: 42px;

            margin-bottom: 22px;
        }

        .login-brand {
            font-size: 22px;
            font-weight: 800;

            letter-spacing: 1px;

            margin-bottom: 15px;
        }

        .login-heading {
            font-size: 50px;
            line-height: 1.06;

            font-weight: 850;

            margin-bottom: 25px;
        }

        .login-description {
            font-size: 17px;
            line-height: 1.75;

            color:
                rgba(255,255,255,0.95);

            max-width: 570px;

            margin-bottom: 30px;
        }

        .feature-row {
            display: flex;
            flex-wrap: wrap;

            gap: 10px;
        }

        .feature-chip {
            padding: 10px 16px;

            border-radius: 999px;

            background:
                rgba(255,255,255,0.15);

            border:
                1px solid
                rgba(255,255,255,0.32);

            color: white;

            font-size: 14px;
            font-weight: 700;
        }

        /* ====================================================
           RIGHT CONTENT
           ==================================================== */

        .login-right-content {
            position: relative;
            z-index: 5;

            padding: 90px 65px 35px 65px;
        }

        .welcome-title {
            font-size: 42px;

            line-height: 1.15;

            font-weight: 850;

            color: #172033;

            margin-bottom: 14px;
        }

        .welcome-text {
            font-size: 18px;

            line-height: 1.7;

            color: #64748B;

            margin-bottom: 30px;
        }

        .welcome-line {
            height: 1px;

            background: #E2E8F0;

            margin-bottom: 30px;
        }

        .input-heading {
            font-size: 16px;

            font-weight: 750;

            color: #172033;

            margin-bottom: 8px;
        }

        /* ====================================================
           INPUT
           ==================================================== */

        div[data-testid="stTextInput"] {
    width: 520px !important;
    max-width: 90% !important;

    margin: 0 auto 18px auto !important;
}

        div[data-testid="stTextInput"] input {
    height: 50px !important;
    min-height: 50px !important;

    border-radius: 12px !important;

    border: 1px solid #CBD5E1 !important;

    background: #F8FAFC !important;

    color: #172033 !important;

    font-size: 15px !important;

    padding: 0 16px !important;

    box-sizing: border-box !important;
}

        div[data-testid="stTextInput"] input:focus {
            border-color:
                #2563EB !important;

            box-shadow:
                0 0 0 3px
                rgba(37,99,235,0.12) !important;
        }

        /* ====================================================
           BUTTON
           ==================================================== */

        .login-button {
    width: 520px !important;
    max-width: 90% !important;

    margin: 0 auto !important;
}

    .login-button button {
    width: 100% !important;

    height: 50px !important;
    min-height: 50px !important;

    border-radius: 12px !important;

    background: #38BDF8 !important;
    color: white !important;

    border: none !important;

    font-size: 15px !important;

    font-weight: 700 !important;

    box-shadow:
        0 6px 15px
        rgba(56, 189, 248, 0.25) !important;
}
.login-button {
    display: flex !important;
    justify-content: center !important;
}

.login-button > div {
    width: 520px !important;
    max-width: 90% !important;
}

.login-button button:hover {
  background: #0EA5E9 !important;
    color: white !important;

    transform: translateY(-2px) !important;
}

        /* ====================================================
           FOOTER
           ==================================================== */

        .login-footer {
            text-align: center;

            color: #94A3B8;

            font-size: 13px;

            margin-top: 32px;
        }

        /* ====================================================
           MOBILE
           ==================================================== */

        @media (max-width: 850px) {

            .main .block-container {
                padding:
                    20px !important;
            }

            [data-testid="stHorizontalBlock"] {
                flex-direction:
                    column !important;

                gap:
                    0 !important;
            }

            [data-testid="stColumn"] {
                min-height:
                    auto !important;
            }

            [data-testid="stColumn"]:first-child {
                border-radius:
                    25px 25px 0 0 !important;
            }

            [data-testid="stColumn"]:last-child {
                border-radius:
                    0 0 25px 25px !important;
            }

            .login-left-content {
                padding: 40px;
            }

            .login-right-content {
                padding: 50px 40px 30px 40px;
            }

            .login-heading {
                font-size: 38px;
            }

            .welcome-title {
                font-size: 34px;
            }
        }

        </style>
        """,
        unsafe_allow_html=True
    )


    # ========================================================
    # TWO JOINED CARDS
    # ========================================================

    left_col, right_col = st.columns(
        [1, 1],
        gap=None
    )


    # ========================================================
    # LEFT BLUE CARD
    # ========================================================

    with left_col:

        render_html(
            """
            <div class="login-left-content">

                <div class="login-logo">
                    📡
                </div>

                <div class="login-brand">
                    CNW AI
                </div>

                <div class="login-heading">
                    Learn.<br>
                    Practice.<br>
                    Perform.
                </div>

                <div class="login-description">
                    Your intelligent learning assistant
                    for Cellular and Wireless Networks.
                    Learn concepts, prepare for exams,
                    practice viva questions and analyze
                    network performance.
                </div>

                <div class="feature-row">

                    <div class="feature-chip">
                        📚 Study
                    </div>

                    <div class="feature-chip">
                        📝 Exam
                    </div>

                    <div class="feature-chip">
                        🎤 Viva
                    </div>

                    <div class="feature-chip">
                        🎯 Quiz
                    </div>

                    <div class="feature-chip">
                        📈 ML Predictor
                    </div>

                </div>

            </div>
            """
        )


    # ========================================================
    # RIGHT WHITE CARD
    # ========================================================

    with right_col:

        render_html(
            """
            <div class="login-right-content">

                <div class="welcome-title">
                    Welcome! 👋
                </div>

                <div class="welcome-text">
                    Start your CNW learning journey.<br>
                    Enter your name below to continue.
                </div>

                <div class="welcome-line"></div>

                

            </div>
            """
        )


        # Name input
        name = st.text_input(
            "",
            placeholder="Enter your name",
            key="login_name",
            label_visibility="collapsed"
        )


        # Start Learning button
        st.markdown(
            '<div class="login-button">',
            unsafe_allow_html=True
        )


        if st.button(
            "🚀  Start Learning",
            use_container_width=True,
            key="login_button"
        ):

            if name.strip():

                st.session_state.logged_in = True

                st.session_state.user_name = name.strip()

                st.session_state.selected_page = "🏠 Home"

                st.rerun()

            else:

                st.warning(
                    "Please enter your name."
                )


        st.markdown(
            "</div>",
            unsafe_allow_html=True
        )


        render_html(
            """
            <div class="login-footer">
                AI-powered Cellular &amp; Wireless
                Network Learning Assistant
            </div>
            """
        )


    st.stop()

# SIDEBAR
# ============================================================

with st.sidebar:

    render_html(
        """
        <div style="
            text-align:center;
            padding:18px 5px 25px 5px;
        ">

            <div style="
                font-size:55px;
                margin-bottom:8px;
            ">
                📡
            </div>

            <div style="
                font-size:30px;
                font-weight:850;
                color:#172033;
            ">
                CNW AI
            </div>

            <div style="
                color:#64748B;
                font-size:16px;
                margin-top:5px;
            ">
                Learning Assistant
            </div>

        </div>

        <hr style="
            border:none;
            border-top:1px solid #DCE6F2;
            margin:5px 0 20px 0;
        ">
        """
    )

    # --------------------------------------------------------
    # HOME
    # --------------------------------------------------------

    if st.button(
        "🏠 Home",
        use_container_width=True,
        key="nav_home"
    ):
        st.session_state.selected_page = "🏠 Home"
        st.rerun()

    st.markdown(
        """
        <div style="
            color:#64748B;
            font-size:14px;
            font-weight:800;
            margin:25px 0 10px 3px;
            letter-spacing:0.5px;
        ">
            LEARN
        </div>
        """,
        unsafe_allow_html=True
    )

    if st.button(
        "📚 Study",
        use_container_width=True,
        key="nav_study"
    ):
        st.session_state.selected_page = "📚 Study"
        st.rerun()

    if st.button(
        "📝 Exam",
        use_container_width=True,
        key="nav_exam"
    ):
        st.session_state.selected_page = "📝 Exam"
        st.rerun()

    if st.button(
        "🎤 Viva",
        use_container_width=True,
        key="nav_viva"
    ):
        st.session_state.selected_page = "🎤 Viva"
        st.rerun()

    st.markdown(
        """
        <div style="
            color:#64748B;
            font-size:14px;
            font-weight:800;
            margin:25px 0 10px 3px;
            letter-spacing:0.5px;
        ">
            INTERACTIVE
        </div>
        """,
        unsafe_allow_html=True
    )

    if st.button(
        "📊 Presentation",
        use_container_width=True,
        key="nav_presentation"
    ):
        st.session_state.selected_page = "📊 Presentation"
        st.rerun()

    if st.button(
        "🧠 Mind Map",
        use_container_width=True,
        key="nav_mindmap"
    ):
        st.session_state.selected_page = "🧠 Mind Map"
        st.rerun()

    if st.button(
        "🎴 Flashcards",
        use_container_width=True,
        key="nav_flashcards"
    ):
        st.session_state.selected_page = "🎴 Flashcards"
        st.rerun()

    if st.button(
        "🎯 Quiz",
        use_container_width=True,
        key="nav_quiz"
    ):
        st.session_state.selected_page = "🎯 Quiz"
        st.rerun()

    st.markdown(
        """
        <div style="
            color:#64748B;
            font-size:14px;
            font-weight:800;
            margin:25px 0 10px 3px;
            letter-spacing:0.5px;
        ">
            ANALYSIS
        </div>
        """,
        unsafe_allow_html=True
    )

    if st.button(
        "📈 Network Predictor",
        use_container_width=True,
        key="nav_predictor"
    ):
        st.session_state.selected_page = "📈 Network Predictor"
        st.rerun()

    st.markdown(
        """
        <div style="
            color:#64748B;
            font-size:14px;
            font-weight:800;
            margin:25px 0 10px 3px;
            letter-spacing:0.5px;
        ">
            PERSONAL
        </div>
        """,
        unsafe_allow_html=True
    )

    if st.button(
        "🕘 History",
        use_container_width=True,
        key="nav_history"
    ):
        st.session_state.selected_page = "🕘 History"
        st.rerun()

    st.markdown(
        """
        <div style="
            color:#64748B;
            font-size:14px;
            font-weight:800;
            margin:25px 0 10px 3px;
            letter-spacing:0.5px;
        ">
            INFO
        </div>
        """,
        unsafe_allow_html=True
    )

    if st.button(
        "ℹ️ About",
        use_container_width=True,
        key="nav_about"
    ):
        st.session_state.selected_page = "ℹ️ About"
        st.rerun()

    if st.button(
        "👥 Team",
        use_container_width=True,
        key="nav_team"
    ):
        st.session_state.selected_page = "👥 Team"
        st.rerun()

    st.markdown(
        "<br>",
        unsafe_allow_html=True
    )

    if st.button(
        "🚪 Sign Out",
        use_container_width=True,
        key="logout"
    ):

        st.session_state.logged_in = False
        st.session_state.user_name = ""

        st.rerun()


# ============================================================
# CURRENT PAGE
# ============================================================

selected_page = st.session_state.selected_page


# ============================================================
# HOME
# ============================================================

if selected_page == "🏠 Home":

    render_html(
        f"""
        <div style="
            margin-bottom:10px;
        ">

            <div style="
                font-size:42px;
                font-weight:850;
                color:#172033;
            ">
                👋 Welcome, {st.session_state.user_name}!
            </div>

            <div style="
                font-size:21px;
                color:#64748B;
                margin-top:8px;
            ">
                Choose a mode to learn, practice, create and analyze.
            </div>

        </div>
        """
    )

    st.write("")

    # ========================================================
    # ROW 1
    # ========================================================

    col1, col2, col3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with col1:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    📚
                </div>

                <div class="mode-title">
                    Study Mode
                </div>

                <div class="mode-description">
                    Understand CNW concepts in simple English
                    with clear AI explanations.
                </div>

            </div>
            """
        )

        if st.button(
            "📚 Open Study Mode",
            use_container_width=True,
            key="home_study"
        ):
            st.session_state.selected_page = "📚 Study"
            st.rerun()

    with col2:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    📝
                </div>

                <div class="mode-title">
                    Exam Mode
                </div>

                <div class="mode-description">
                    Get structured 2, 5 and 10-mark
                    examination answers.
                </div>

            </div>
            """
        )

        if st.button(
            "📝 Open Exam Mode",
            use_container_width=True,
            key="home_exam"
        ):
            st.session_state.selected_page = "📝 Exam"
            st.rerun()

    with col3:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    🎤
                </div>

                <div class="mode-title">
                    Viva Mode
                </div>

                <div class="mode-description">
                    Practice viva questions and receive
                    instant feedback.
                </div>

            </div>
            """
        )

        if st.button(
            "🎤 Open Viva Mode",
            use_container_width=True,
            key="home_viva"
        ):
            st.session_state.selected_page = "🎤 Viva"
            st.rerun()

    # ========================================================
    # ROW 2
    # ========================================================

    col1, col2, col3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with col1:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    📊
                </div>

                <div class="mode-title">
                    Presentation
                </div>

                <div class="mode-description">
                    Create a colorful PowerPoint presentation
                    from any CNW topic.
                </div>

            </div>
            """
        )

        if st.button(
            "📊 Open Presentation",
            use_container_width=True,
            key="home_presentation"
        ):
            st.session_state.selected_page = "📊 Presentation"
            st.rerun()

    with col2:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    🧠
                </div>

                <div class="mode-title">
                    Mind Map
                </div>

                <div class="mode-description">
                    Turn a CNW topic into an organized
                    revision structure.
                </div>

            </div>
            """
        )

        if st.button(
            "🧠 Open Mind Map",
            use_container_width=True,
            key="home_mindmap"
        ):
            st.session_state.selected_page = "🧠 Mind Map"
            st.rerun()

    with col3:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    🎴
                </div>

                <div class="mode-title">
                    Flashcards
                </div>

                <div class="mode-description">
                    Revise important CNW concepts using
                    colorful interactive cards.
                </div>

            </div>
            """
        )

        if st.button(
            "🎴 Open Flashcards",
            use_container_width=True,
            key="home_flashcards"
        ):
            st.session_state.selected_page = "🎴 Flashcards"
            st.rerun()

    # ========================================================
    # ROW 3
    # ========================================================

    col1, col2, col3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with col1:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    🎯
                </div>

                <div class="mode-title">
                    Quiz
                </div>

                <div class="mode-description">
                    Test your CNW knowledge with interactive
                    multiple-choice questions.
                </div>

            </div>
            """
        )

        if st.button(
            "🎯 Open Quiz",
            use_container_width=True,
            key="home_quiz"
        ):
            st.session_state.selected_page = "🎯 Quiz"
            st.rerun()

    with col2:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    📈
                </div>

                <div class="mode-title">
                    Network Predictor
                </div>

                <div class="mode-description">
                    Analyze cellular network performance
                    using important KPIs.
                </div>

            </div>
            """
        )

        if st.button(
            "📈 Open Predictor",
            use_container_width=True,
            key="home_predictor"
        ):
            st.session_state.selected_page = "📈 Network Predictor"
            st.rerun()

    with col3:

        render_html(
            """
            <div class="mode-card">

                <div class="mode-icon">
                    🕘
                </div>

                <div class="mode-title">
                    History
                </div>

                <div class="mode-description">
                    View your recently used learning
                    topics and activities.
                </div>

            </div>
            """
        )

        if st.button(
            "🕘 Open History",
            use_container_width=True,
            key="home_history"
        ):
            st.session_state.selected_page = "🕘 History"
            st.rerun()


# ============================================================
# STUDY MODE
# ============================================================

elif selected_page == "📚 Study":

    st.markdown('<div class="page-title">📚 Study Mode</div>', unsafe_allow_html=True)
    st.markdown('<div class="page-subtitle">Understand Cellular and Wireless Network concepts in simple English.</div>', unsafe_allow_html=True)

    # Generate any pending question
    if st.session_state.study_pending_question:
        pending = st.session_state.study_pending_question
        st.session_state.study_pending_question = ""

        with st.spinner("✨ Preparing your explanation..."):
            prompt = f"""You are a Cellular and Wireless Networks teacher.
Answer this question in simple English:

{pending}

Requirements:
- Give a clear and accurate explanation.
- Use simple English.
- Use headings or bullet points where useful.
- Give an example where appropriate.
- Keep the answer focused.
- Do not use Markdown headings, **, ***, ###, ``` or unnecessary Markdown symbols."""
            answer = clean_ai_text(ask_gemini(prompt))

        st.session_state.study_chat.append({"question": pending, "answer": answer})
        add_history("Study", pending)

    # Conversation history
    for item in st.session_state.study_chat:
        render_html(f"""
        <div style="background:#EFF6FF;border-radius:16px;padding:18px 22px;margin:18px 0 10px 0;border:1px solid #BFDBFE;">
            <div style="font-weight:800;color:#2563EB;margin-bottom:8px;">👤 You</div>
            <div style="font-size:18px;color:#172033;line-height:1.6;">{format_ai_text(item['question'])}</div>
        </div>
        <div class="answer-card">
            <div class="answer-title">🤖 CNW AI</div>
            <div class="answer-content">{format_ai_text(item['answer'])}</div>
        </div>
        """)

    # New question at the bottom
    # Use a fresh widget key for every turn. This avoids modifying a Streamlit
    # widget's session-state value after that widget has already been created.
    study_turn = len(st.session_state.study_chat) + 1
    study_text_key = f"study_chat_input_{study_turn}"
    study_voice_key = f"study_chat_voice_{study_turn}"
    study_send_key = f"study_chat_send_{study_turn}"

    study_question, submitted = chat_input_with_voice(
        study_text_key,
        study_voice_key,
        "Ask your next CNW question...",
        study_send_key,
        "✨ Ask"
    )

    if submitted:
        if study_question.strip():
            st.session_state.study_pending_question = study_question.strip()
            st.rerun()
        else:
            st.warning("Please enter a question.")


# ============================================================
# EXAM MODE
# ============================================================

elif selected_page == "📝 Exam":

    st.markdown('<div class="page-title">📝 Exam Mode</div>', unsafe_allow_html=True)
    st.markdown('<div class="page-subtitle">Get structured 2, 5 or 10-mark answers for your CNW exam.</div>', unsafe_allow_html=True)

    exam_mark = st.selectbox(
        "Select answer type",
        ["2 Marks", "5 Marks", "10 Marks"],
        key="exam_mark_selector"
    )
    st.session_state.exam_mark = exam_mark

    if st.session_state.exam_pending_question:
        pending = st.session_state.exam_pending_question
        st.session_state.exam_pending_question = ""

        with st.spinner("✨ Preparing your exam answer..."):
            prompt = f"""You are an expert Cellular and Wireless Networks exam-answer assistant.

Question:
{pending}

Required answer length: {exam_mark}

For 2 Marks: Give a direct definition and the most important point.
For 5 Marks: Give definition, explanation, important points and an example where useful.
For 10 Marks: Give introduction/definition, detailed explanation, important points, advantages/disadvantages/applications where relevant, and a suitable example or conclusion.

Use simple English and make the answer suitable for an engineering examination. Do not use Markdown headings, **, ***, ###, ``` or unnecessary Markdown symbols."""
            answer = clean_ai_text(ask_gemini(prompt))

        st.session_state.exam_chat.append({
            "question": pending,
            "answer": answer,
            "marks": exam_mark
        })
        add_history("Exam", f"{exam_mark}: {pending}")

    # Conversation history
    for item in st.session_state.exam_chat:
        render_html(f"""
        <div style="background:#EFF6FF;border-radius:16px;padding:18px 22px;margin:18px 0 10px 0;border:1px solid #BFDBFE;">
            <div style="font-weight:800;color:#2563EB;margin-bottom:8px;">👤 You</div>
            <div style="font-size:18px;color:#172033;line-height:1.6;">{format_ai_text(item['question'])}</div>
        </div>
        <div class="answer-card">
            <div class="answer-title">📝 CNW AI • {item['marks']}</div>
            <div class="answer-content">{format_ai_text(item['answer'])}</div>
        </div>
        """)

    # Use a fresh widget key for every turn so the previous question
    # naturally disappears from the new input while the chat history remains.
    exam_turn = len(st.session_state.exam_chat) + 1
    exam_text_key = f"exam_chat_input_{exam_turn}"
    exam_voice_key = f"exam_chat_voice_{exam_turn}"
    exam_send_key = f"exam_chat_send_{exam_turn}"

    exam_question, submitted = chat_input_with_voice(
        exam_text_key,
        exam_voice_key,
        "Ask your next exam question...",
        exam_send_key,
        "📝 Ask"
    )

    if submitted:
        if exam_question.strip():
            st.session_state.exam_pending_question = exam_question.strip()
            st.rerun()
        else:
            st.warning("Please enter an exam question.")


# ============================================================
# VIVA MODE
# ============================================================

elif selected_page == "🎤 Viva":

    st.markdown(
        '<div class="page-title">🎤 Viva Mode</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Practice viva questions and receive feedback on your answers.'
        '</div>',
        unsafe_allow_html=True
    )

    if not st.session_state.viva_started:

        viva_topic = st.text_input(
            "Enter viva topic",
            placeholder="Example: OFDM, MIMO, LTE, 5G"
        )

        if st.button(
            "🎤 Start Viva",
            use_container_width=True
        ):

            if viva_topic.strip():

                prompt = f"""
Create exactly 5 viva questions about:

{viva_topic}

Return ONLY valid JSON:

[
  {{
    "question": "Question"
  }}
]

Requirements:
- Technically accurate.
- Suitable for an engineering viva.
- Questions should progress from basic to moderate difficulty.
- Simple English.
"""

                response = ask_gemini(prompt)

                try:

                    questions = json.loads(
                        clean_json_text(response)
                    )

                    st.session_state.viva_questions = questions
                    st.session_state.viva_index = 0
                    st.session_state.viva_score = 0
                    st.session_state.viva_answered = False
                    st.session_state.viva_feedback = ""
                    st.session_state.viva_started = True

                    add_history(
                        "Viva",
                        viva_topic
                    )

                    st.rerun()

                except Exception as e:

                    st.error(
                        f"Could not create viva: {e}"
                    )

            else:

                st.warning(
                    "Please enter a topic."
                )

    else:

        questions = st.session_state.viva_questions
        index = st.session_state.viva_index

        if index >= len(questions):

            render_html(
                f"""
                <div style="
                    background:
                    linear-gradient(
                        135deg,
                        #2563EB,
                        #06B6D4
                    );
                    color:white;
                    border-radius:28px;
                    padding:45px;
                    text-align:center;
                ">

                    <div style="
                        font-size:55px;
                    ">
                        🎉
                    </div>

                    <h1>
                        Viva Complete!
                    </h1>

                    <h2>
                        Score:
                        {st.session_state.viva_score}
                        /
                        {len(questions)}
                    </h2>

                </div>
                """
            )

            st.write("")

            if st.button(
                "🔄 Start New Viva",
                use_container_width=True
            ):

                st.session_state.viva_started = False
                st.session_state.viva_questions = []
                st.session_state.viva_index = 0
                st.session_state.viva_score = 0
                st.session_state.viva_answered = False
                st.session_state.viva_feedback = ""

                st.rerun()

        else:

            q = questions[index]

            render_html(
                f"""
                <div class="info-card"
                     style="
                     border-top:6px solid #06B6D4;
                     ">

                    <div style="
                        color:#0891B2;
                        font-weight:800;
                        margin-bottom:10px;
                    ">
                        VIVA QUESTION
                        {index + 1}
                        /
                        {len(questions)}
                    </div>

                    <div style="
                        font-size:25px;
                        font-weight:750;
                        color:#172033;
                    ">
                        {q.get("question", "")}
                    </div>

                </div>
                """
            )

            viva_answer = st.text_area(
                "Your answer",
                placeholder="Type your viva answer here...",
                height=150,
                key=f"viva_answer_{index}",
                disabled=st.session_state.viva_answered
            )

            if not st.session_state.viva_answered:

                if st.button(
                    "✅ Submit Answer",
                    use_container_width=True
                ):

                    if viva_answer.strip():

                        with st.spinner(
                            "🎤 Evaluating your answer..."
                        ):

                            prompt = f"""
You are a Cellular and Wireless Networks viva examiner.

Question:
{q.get("question", "")}

Student answer:
{viva_answer}

Evaluate the answer.

Return ONLY valid JSON:

{{
  "correct": true,
  "feedback": "Short feedback",
  "ideal_answer": "Short ideal answer"
}}

Use true or false for correct.
Keep feedback concise.
"""

                            response = ask_gemini(prompt)

                        try:

                            feedback_data = json.loads(
                                clean_json_text(response)
                            )

                            if feedback_data.get(
                                "correct",
                                False
                            ):

                                st.session_state.viva_score += 1

                            st.session_state.viva_feedback = (
                                feedback_data
                            )

                            st.session_state.viva_answered = True

                            st.rerun()

                        except Exception as e:

                            st.error(
                                f"Could not evaluate answer: {e}"
                            )

                    else:

                        st.warning(
                            "Please enter your answer."
                        )

            else:

                feedback = st.session_state.viva_feedback

                if feedback.get("correct", False):

                    render_html(
                        """
                        <div style="
                            background:#ECFDF5;
                            border:2px solid #10B981;
                            border-radius:18px;
                            padding:22px;
                            margin-top:15px;
                        ">

                            <h2 style="
                                color:#047857;
                                margin-top:0;
                            ">
                                🎉 Correct
                            </h2>

                        </div>
                        """
                    )

                else:

                    render_html(
                        """
                        <div style="
                            background:#FEF2F2;
                            border:2px solid #EF4444;
                            border-radius:18px;
                            padding:22px;
                            margin-top:15px;
                        ">

                            <h2 style="
                                color:#DC2626;
                                margin-top:0;
                            ">
                                ❌ Needs Improvement
                            </h2>

                        </div>
                        """
                    )

                render_html(
                    f"""
                    <div class="answer-card">

                        <div class="answer-title">
                            💬 Feedback
                        </div>

                        <div class="answer-content">
                            {format_ai_text(
                                feedback.get(
                                    "feedback",
                                    ""
                                )
                            )}
                        </div>

                        <br>

                        <div class="answer-title">
                            📖 Ideal Answer
                        </div>

                        <div class="answer-content">
                            {format_ai_text(
                                feedback.get(
                                    "ideal_answer",
                                    ""
                                )
                            )}
                        </div>

                    </div>
                    """
                )

                if st.button(
                    "➡️ Next Viva Question",
                    use_container_width=True
                ):

                    st.session_state.viva_index += 1
                    st.session_state.viva_answered = False
                    st.session_state.viva_feedback = ""

                    st.rerun()


# ============================================================
# PRESENTATION
# ============================================================

elif selected_page == "📊 Presentation":

    st.markdown(
        '<div class="page-title">📊 Presentation</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Create a colorful PowerPoint presentation from any CNW topic.'
        '</div>',
        unsafe_allow_html=True
    )

    if not PPT_AVAILABLE:

        st.warning(
            "python-pptx is not installed. "
            "Run: pip install python-pptx"
        )

    presentation_topic = st.text_input(
        "Enter presentation topic",
        placeholder="Example: 5G Architecture"
    )

    slide_count = st.selectbox(
        "Number of slides",
        [5, 6, 7, 8, 10],
        index=1
    )

    if st.button(
        "📊 Create Presentation",
        use_container_width=True
    ):

        if presentation_topic.strip():

            prompt = f"""
Create a {slide_count}-slide presentation about:

{presentation_topic}

Return ONLY valid JSON:

[
  {{
    "title": "Slide title",
    "points": [
      "Point 1",
      "Point 2",
      "Point 3",
      "Point 4"
    ]
  }}
]

Requirements:
- Technically accurate.
- Simple English.
- Suitable for engineering students.
- Include introduction, important concepts,
  applications and conclusion where appropriate.
"""

            with st.spinner(
                "📊 Creating your presentation..."
            ):

                response = ask_gemini(prompt)

            try:

                data = json.loads(
                    clean_json_text(response)
                )

                st.session_state.presentation_data = data

                add_history(
                    "Presentation",
                    presentation_topic
                )

                st.rerun()

            except Exception as e:

                st.error(
                    f"Could not create presentation: {e}"
                )

        else:

            st.warning(
                "Please enter a topic."
            )

    # --------------------------------------------------------
    # PRESENTATION PREVIEW
    # --------------------------------------------------------

    if st.session_state.presentation_data:

        data = st.session_state.presentation_data

        st.markdown(
            "### 📑 Presentation Preview"
        )

        for i, slide in enumerate(data):

            render_html(
                f"""
                <div style="
                    background:
                    linear-gradient(
                        135deg,
                        #EFF6FF,
                        #ECFEFF
                    );
                    border:1px solid #BFDBFE;
                    border-radius:20px;
                    padding:25px;
                    margin:15px 0;
                ">

                    <div style="
                        color:#2563EB;
                        font-weight:800;
                        font-size:14px;
                        margin-bottom:8px;
                    ">
                        SLIDE {i + 1}
                    </div>

                    <div style="
                        font-size:25px;
                        font-weight:800;
                        color:#172033;
                        margin-bottom:15px;
                    ">
                        {slide.get("title", "")}
                    </div>

                </div>
                """
            )

            for point in slide.get("points", []):

                st.markdown(
                    f"- {point}"
                )

        # ----------------------------------------------------
        # CREATE PPTX
        # ----------------------------------------------------

        if PPT_AVAILABLE:

            if st.button(
                "⬇️ Generate PowerPoint",
                use_container_width=True
            ):

                prs = Presentation()

                # Title slide
                title_slide = prs.slides.add_slide(
                    prs.slide_layouts[6]
                )

                background = title_slide.background
                fill = background.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor(
                    239, 246, 255
                )

                textbox = title_slide.shapes.add_textbox(
                    Inches(1),
                    Inches(2),
                    Inches(11),
                    Inches(2)
                )

                tf = textbox.text_frame
                tf.clear()

                p = tf.paragraphs[0]
                p.text = presentation_topic
                p.font.size = Pt(34)
                p.font.bold = True
                p.font.color.rgb = RGBColor(
                    37, 99, 235
                )
                p.alignment = PP_ALIGN.CENTER

                # Content slides
                for slide_data in data:

                    slide = prs.slides.add_slide(
                        prs.slide_layouts[6]
                    )

                    background = slide.background
                    fill = background.fill
                    fill.solid()
                    fill.fore_color.rgb = RGBColor(
                        248, 250, 252
                    )

                    # Header
                    header = slide.shapes.add_shape(
                        1,
                        Inches(0),
                        Inches(0),
                        Inches(13.333),
                        Inches(1.2)
                    )

                    header.fill.solid()
                    header.fill.fore_color.rgb = RGBColor(
                        239, 246, 255
                    )

                    header.line.color.rgb = RGBColor(
                        191, 219, 254
                    )

                    title_box = slide.shapes.add_textbox(
                        Inches(0.7),
                        Inches(0.3),
                        Inches(11.8),
                        Inches(0.7)
                    )

                    tf = title_box.text_frame
                    tf.clear()

                    p = tf.paragraphs[0]
                    p.text = slide_data.get(
                        "title",
                        ""
                    )
                    p.font.size = Pt(27)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(
                        37, 99, 235
                    )

                    # Content
                    content_box = slide.shapes.add_textbox(
                        Inches(0.9),
                        Inches(1.7),
                        Inches(11.4),
                        Inches(4.8)
                    )

                    tf = content_box.text_frame
                    tf.clear()

                    for j, point in enumerate(
                        slide_data.get("points", [])
                    ):

                        if j == 0:
                            p = tf.paragraphs[0]
                        else:
                            p = tf.add_paragraph()

                        p.text = "• " + str(point)
                        p.font.size = Pt(21)
                        p.font.color.rgb = RGBColor(
                            51, 65, 85
                        )
                        p.space_after = Pt(15)

                ppt_path = "CNW_AI_Presentation.pptx"

                prs.save(ppt_path)

                with open(
                    ppt_path,
                    "rb"
                ) as file:

                    st.download_button(
                        "📥 Download PowerPoint",
                        data=file,
                        file_name=ppt_path,
                        mime=(
                            "application/vnd.openxmlformats-"
                            "officedocument.presentationml.presentation"
                        ),
                        use_container_width=True
                    )


# ============================================================
# MIND MAP
# ============================================================

elif selected_page == "🧠 Mind Map":

    st.markdown(
        '<div class="page-title">🧠 Mind Map</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Organize a CNW topic into an easy-to-revise structure.'
        '</div>',
        unsafe_allow_html=True
    )

    mind_topic = st.text_input(
        "Enter topic",
        placeholder="Example: SISO"
    )

    if st.button(
        "🧠 Create Mind Map",
        use_container_width=True,
        key="create_mindmap"
    ):

        if not mind_topic.strip():

            st.warning("Please enter a topic.")

        else:

            prompt = f"""
Create a clear mind map for the Cellular and Wireless Networks topic:

{mind_topic}

Return ONLY valid JSON.

The response MUST be a single JSON OBJECT, not a JSON list.

Use EXACTLY this structure:

{{
    "central_topic": "{mind_topic}",
    "branches": [
        {{
            "title": "Branch 1",
            "points": [
                "Point 1",
                "Point 2",
                "Point 3"
            ]
        }},
        {{
            "title": "Branch 2",
            "points": [
                "Point 1",
                "Point 2",
                "Point 3"
            ]
        }}
    ]
}}

Requirements:

- Return one JSON object only.
- Do NOT return a list.
- Do NOT use Markdown.
- Do NOT use ```json.
- Create 5 to 7 branches.
- Each branch should contain 3 to 5 important points.
- Use simple English.
- Keep the information technically accurate.
- Focus on important exam and revision points.
"""

            with st.spinner(
                "🧠 Opening Mind Map and preparing your structure..."
            ):

                response = ask_gemini(prompt)

            try:

                cleaned = clean_json_text(response)

                data = json.loads(cleaned)

                # ------------------------------------------------
                # HANDLE GEMINI RETURNING A LIST
                # ------------------------------------------------

                if isinstance(data, list):

                    # Sometimes Gemini returns only the branches.
                    # Convert that list into the expected structure.

                    data = {
                        "central_topic": mind_topic.strip(),
                        "branches": data
                    }

                elif not isinstance(data, dict):

                    raise ValueError(
                        "Gemini returned an unsupported mind-map format."
                    )

                # ------------------------------------------------
                # GET CENTRAL TOPIC
                # ------------------------------------------------

                central_topic = data.get(
                    "central_topic",
                    mind_topic.strip()
                )

                # ------------------------------------------------
                # GET BRANCHES
                # ------------------------------------------------

                branches = data.get(
                    "branches",
                    []
                )

                if not isinstance(branches, list):

                    branches = []

                # ------------------------------------------------
                # NORMALIZE BRANCHES
                # ------------------------------------------------

                clean_branches = []

                for branch in branches:

                    # If Gemini returns a dictionary
                    if isinstance(branch, dict):

                        title = branch.get(
                            "title",
                            "Important Concept"
                        )

                        points = branch.get(
                            "points",
                            []
                        )

                        if not isinstance(points, list):
                            points = [str(points)]

                        points = [
                            str(point)
                            for point in points
                            if str(point).strip()
                        ]

                        clean_branches.append(
                            {
                                "title": str(title),
                                "points": points
                            }
                        )

                    # If Gemini returns a simple string
                    elif isinstance(branch, str):

                        clean_branches.append(
                            {
                                "title": branch,
                                "points": []
                            }
                        )

                if not clean_branches:

                    raise ValueError(
                        "No valid branches were returned. Please try again."
                    )

                # ------------------------------------------------
                # SAVE MIND MAP
                # ------------------------------------------------

                st.session_state.mindmap_data = {
                    "central_topic": central_topic,
                    "branches": clean_branches
                }

                add_history(
                    "Mind Map",
                    mind_topic
                )

                st.rerun()

            except Exception as e:

                st.error(
                    f"Could not create mind map: {e}"
                )

                st.code(
                    str(response),
                    language="text"
                )

    # ============================================================
    # DISPLAY MIND MAP
    # ============================================================

    if st.session_state.mindmap_data:

        data = st.session_state.mindmap_data

        # --------------------------------------------------------
        # CENTRAL TOPIC
        # --------------------------------------------------------

        central_topic = data.get(
            "central_topic",
            mind_topic if "mind_topic" in locals() else "CNW Topic"
        )

        render_html(
            f"""
            <div style="
                text-align:center;
                background:
                linear-gradient(
                    135deg,
                    #2563EB,
                    #06B6D4
                );
                color:white;
                border-radius:25px;
                padding:30px;
                margin:25px 0;
                box-shadow:
                0 15px 35px
                rgba(37,99,235,0.18);
            ">

                <div style="
                    font-size:15px;
                    font-weight:800;
                    letter-spacing:1px;
                    opacity:0.9;
                ">
                    🧠 CENTRAL TOPIC
                </div>

                <div style="
                    font-size:34px;
                    font-weight:900;
                    margin-top:8px;
                ">
                    {central_topic}
                </div>

            </div>
            """
        )

        # --------------------------------------------------------
        # BRANCHES
        # --------------------------------------------------------

        branches = data.get(
            "branches",
            []
        )

        if isinstance(branches, list):

            for i in range(
                0,
                len(branches),
                2
            ):

                c1, c2 = st.columns(
                    [1, 1],
                    gap="medium"
                )

                current_branches = branches[
                    i:i + 2
                ]

                for col, branch in zip(
                    [c1, c2],
                    current_branches
                ):

                    with col:

                        if not isinstance(
                            branch,
                            dict
                        ):
                            continue

                        title = branch.get(
                            "title",
                            "Important Concept"
                        )

                        points = branch.get(
                            "points",
                            []
                        )

                        if not isinstance(
                            points,
                            list
                        ):
                            points = [
                                str(points)
                            ]

                        # ------------------------------------------------
                        # BRANCH CARD
                        # ------------------------------------------------

                        points_html = ""

                        for point in points:

                            points_html += f"""
                            <div style="
                                background:#F8FAFC;
                                border-left:
                                4px solid #06B6D4;
                                border-radius:10px;
                                padding:10px 12px;
                                margin:8px 0;
                                color:#334155;
                                font-size:16px;
                                line-height:1.5;
                            ">
                                • {point}
                            </div>
                            """

                        render_html(
                            f"""
                            <div style="
                                background:#FFFFFF;
                                border:1px solid #DCE6F2;
                                border-top:
                                6px solid #2563EB;
                                border-radius:20px;
                                padding:22px;
                                margin-bottom:20px;
                                box-shadow:
                                0 8px 22px
                                rgba(37,99,235,0.08);
                            ">

                                <div style="
                                    color:#2563EB;
                                    font-size:22px;
                                    font-weight:850;
                                    margin-bottom:14px;
                                ">
                                    🧩 {title}
                                </div>

                                {points_html}

                            </div>
                            """
                        )

# ============================================================
# FLASHCARDS
# ============================================================

elif selected_page == "🎴 Flashcards":

    st.markdown(
        '<div class="page-title">🎴 Flashcards</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Revise important CNW concepts using colorful cards.'
        '</div>',
        unsafe_allow_html=True
    )

    if not st.session_state.flashcards:

        flash_topic = st.text_input(
            "Enter topic",
            placeholder="Example: LTE"
        )

        if st.button(
            "🎴 Create Flashcards",
            use_container_width=True
        ):

            if flash_topic.strip():

                prompt = f"""
Create 6 flashcards about:

{flash_topic}

Return ONLY valid JSON:

[
  {{
    "title": "Concept",
    "content": "Complete explanation"
  }}
]

Requirements:
- Every card must contain all necessary information.
- Do not require a Show Answer button.
- Use simple English.
"""

                with st.spinner(
                    "🎴 Creating flashcards..."
                ):

                    response = ask_gemini(prompt)

                try:

                    cards = json.loads(
                        clean_json_text(response)
                    )

                    st.session_state.flashcards = cards
                    st.session_state.flashcard_index = 0

                    add_history(
                        "Flashcards",
                        flash_topic
                    )

                    st.rerun()

                except Exception as e:

                    st.error(
                        f"Could not create flashcards: {e}"
                    )

    else:

        cards = st.session_state.flashcards
        index = st.session_state.flashcard_index

        if cards:

            card = cards[index]

            render_html(
                f"""
                <div class="flashcard">

                    <div class="flashcard-number">
                        CARD {index + 1} / {len(cards)}
                    </div>

                    <div class="flashcard-title">
                        {card.get("title", "")}
                    </div>

                    <div class="flashcard-content">
                        {format_ai_text(
                            card.get(
                                "content",
                                ""
                            )
                        )}
                    </div>

                </div>
                """
            )

            st.write("")

            c1, c2 = st.columns(
                [1, 1],
                gap="medium"
            )

            with c1:

                if st.button(
                    "⬅️ Previous",
                    use_container_width=True
                ):

                    if index > 0:

                        st.session_state.flashcard_index -= 1
                        st.rerun()

            with c2:

                if st.button(
                    "Next ➡️",
                    use_container_width=True
                ):

                    if index < len(cards) - 1:

                        st.session_state.flashcard_index += 1
                        st.rerun()

            st.write("")

            if st.button(
                "🔄 Create New Flashcards",
                use_container_width=True
            ):

                st.session_state.flashcards = []
                st.session_state.flashcard_index = 0

                st.rerun()


# ============================================================
# QUIZ
# ============================================================

elif selected_page == "🎯 Quiz":

    st.markdown(
        '<div class="page-title">🎯 Quiz</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Test your CNW knowledge with an interactive quiz.'
        '</div>',
        unsafe_allow_html=True
    )

    if not st.session_state.quiz_started:

        quiz_topic = st.text_input(
            "Enter quiz topic",
            placeholder="Example: Cellular Network Architecture"
        )

        if st.button(
            "🎯 Start Quiz",
            use_container_width=True
        ):

            if quiz_topic.strip():

                prompt = f"""
Create exactly 5 multiple-choice questions about:

{quiz_topic}

Return ONLY valid JSON:

[
 {{
   "question": "Question",
   "options": [
      "Option A",
      "Option B",
      "Option C",
      "Option D"
   ],
   "answer": "Correct option",
   "explanation": "Short explanation"
 }}
]

Requirements:
- Questions must be technically accurate.
- Use simple English.
- Each question must have exactly 4 options.
- The answer must exactly match one option.
"""

                with st.spinner(
                    "🎯 Creating quiz..."
                ):

                    response = ask_gemini(prompt)

                try:

                    questions = json.loads(
                        clean_json_text(response)
                    )

                    st.session_state.quiz_questions = questions
                    st.session_state.quiz_index = 0
                    st.session_state.quiz_score = 0
                    st.session_state.quiz_answered = False
                    st.session_state.quiz_started = True

                    add_history(
                        "Quiz",
                        quiz_topic
                    )

                    st.rerun()

                except Exception as e:

                    st.error(
                        f"Could not create quiz: {e}"
                    )

            else:

                st.warning(
                    "Please enter a topic."
                )

    else:

        questions = st.session_state.quiz_questions
        index = st.session_state.quiz_index

        if index >= len(questions):

            score = st.session_state.quiz_score
            total = len(questions)

            render_html(
                f"""
                <div style="
                    background:
                    linear-gradient(
                        135deg,
                        #2563EB,
                        #06B6D4
                    );
                    color:white;
                    border-radius:28px;
                    padding:45px;
                    text-align:center;
                ">

                    <div style="
                        font-size:55px;
                    ">
                        🎉
                    </div>

                    <h1>
                        Quiz Complete!
                    </h1>

                    <h2>
                        Your Score:
                        {score} / {total}
                    </h2>

                    <p style="
                        font-size:18px;
                    ">
                        Great job! Keep practicing CNW concepts.
                    </p>

                </div>
                """
            )

            st.write("")

            if st.button(
                "🔄 Start New Quiz",
                use_container_width=True
            ):

                st.session_state.quiz_started = False
                st.session_state.quiz_questions = []
                st.session_state.quiz_index = 0
                st.session_state.quiz_score = 0
                st.session_state.quiz_answered = False

                st.rerun()

        else:

            q = questions[index]

            render_html(
                f"""
                <div class="info-card"
                     style="
                     border-top:6px solid #2563EB;
                     ">

                    <div style="
                        color:#2563EB;
                        font-weight:800;
                        font-size:15px;
                        margin-bottom:10px;
                    ">
                        QUESTION
                        {index + 1}
                        /
                        {len(questions)}
                    </div>

                    <div style="
                        color:#172033;
                        font-size:25px;
                        font-weight:750;
                    ">
                        {q.get("question", "")}
                    </div>

                </div>
                """
            )

            options = q.get(
                "options",
                []
            )

            selected_answer = st.radio(
                "Choose your answer:",
                options,
                key=f"quiz_option_{index}",
                disabled=st.session_state.quiz_answered
            )

            if not st.session_state.quiz_answered:

                if st.button(
                    "✅ Submit Answer",
                    use_container_width=True,
                    key=f"submit_{index}"
                ):

                    st.session_state.quiz_answered = True

                    if (
                        selected_answer
                        ==
                        q.get("answer")
                    ):

                        st.session_state.quiz_score += 1

                    st.rerun()

            else:

                correct_answer = q.get(
                    "answer",
                    ""
                )

                explanation = q.get(
                    "explanation",
                    ""
                )

                if selected_answer == correct_answer:

                    render_html(
                        """
                        <div style="
                            background:#ECFDF5;
                            border:2px solid #10B981;
                            border-radius:18px;
                            padding:22px;
                        ">

                            <h2 style="
                                color:#047857;
                                margin-top:0;
                            ">
                                🎉 Correct Answer!
                            </h2>

                            <p>
                                Excellent! Your answer is correct.
                            </p>

                        </div>
                        """
                    )

                else:

                    render_html(
                        f"""
                        <div style="
                            background:#FEF2F2;
                            border:2px solid #EF4444;
                            border-radius:18px;
                            padding:22px;
                        ">

                            <h2 style="
                                color:#DC2626;
                                margin-top:0;
                            ">
                                ❌ Incorrect Answer
                            </h2>

                            <p>
                                <b>Your answer:</b>
                                {selected_answer}
                            </p>

                            <p>
                                <b>Correct answer:</b>
                                {correct_answer}
                            </p>

                        </div>
                        """
                    )

                render_html(
                    f"""
                    <div style="
                        background:#ECFEFF;
                        border:2px solid #06B6D4;
                        border-radius:18px;
                        padding:22px;
                        margin-top:15px;
                    ">

                        <h3 style="
                            color:#0891B2;
                            margin-top:0;
                        ">
                            💡 Explanation
                        </h3>

                        <p style="
                            font-size:17px;
                            line-height:1.6;
                        ">
                            {format_ai_text(
                                explanation
                            )}
                        </p>

                    </div>
                    """
                )

                st.write("")

                if index < len(questions) - 1:

                    if st.button(
                        "➡️ Next Question",
                        use_container_width=True,
                        key=f"next_{index}"
                    ):

                        st.session_state.quiz_index += 1
                        st.session_state.quiz_answered = False

                        st.rerun()

                else:

                    if st.button(
                        "🏁 Finish Quiz",
                        use_container_width=True,
                        key="finish_quiz"
                    ):

                        st.session_state.quiz_index += 1

                        st.rerun()


# ============================================================
# NETWORK PREDICTOR
# ============================================================

elif selected_page == "📈 Network Predictor":

    st.markdown(
        '<div class="page-title">📈 Network Performance Predictor</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Predict cellular network QoS using important network KPIs.'
        '</div>',
        unsafe_allow_html=True
    )

    render_html(
        """
        <div class="info-card"
             style="border-left:6px solid #06B6D4;">

            <div style="
                color:#2563EB;
                font-size:25px;
                font-weight:800;
                margin-bottom:15px;
            ">
                🤖 QoS Prediction
            </div>

            <p>
                Enter the cellular network parameters below.
                The trained Random Forest model will predict
                the expected network QoS category.
            </p>

        </div>
        """
    )

    st.markdown("## 📊 Enter Network KPIs")

    # --------------------------------------------------------
    # ROW 1
    # --------------------------------------------------------

    c1, c2, c3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with c1:

        rsrp = st.number_input(
            "RSRP (dBm)",
            value=-92.0,
            step=1.0,
            key="predict_rsrp"
        )

    with c2:

        rsrq = st.number_input(
            "RSRQ (dB)",
            value=-12.0,
            step=1.0,
            key="predict_rsrq"
        )

    with c3:

        rssi = st.number_input(
            "RSSI (dBm)",
            value=-63.0,
            step=1.0,
            key="predict_rssi"
        )

    # --------------------------------------------------------
    # ROW 2
    # --------------------------------------------------------

    c1, c2, c3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with c1:

        sinr = st.number_input(
            "SINR (dB)",
            value=8.0,
            step=1.0,
            key="predict_sinr"
        )

    with c2:

        delay = st.number_input(
            "Delay (ms)",
            value=85.0,
            step=1.0,
            min_value=0.0,
            key="predict_delay"
        )

    with c3:

        throughput_uplink = st.number_input(
            "Uplink Throughput (Mbps)",
            value=200.0,
            step=10.0,
            min_value=0.0,
            key="predict_uplink"
        )

    # --------------------------------------------------------
    # ROW 3
    # --------------------------------------------------------

    c1, c2 = st.columns(
        [1, 1],
        gap="medium"
    )

    with c1:

        ran = st.selectbox(
            "RAN",
            ["5G-NSA", "LTE"],
            key="predict_ran"
        )

    with c2:

        band = st.selectbox(
            "Band",
            [
                "LTE_B20",
                "LTE_B3",
                "LTE_B7",
                "LTE_B8",
                "LTE_B1"
            ],
            key="predict_band"
        )

    st.markdown("<br>", unsafe_allow_html=True)

    # --------------------------------------------------------
    # PREDICT BUTTON
    # --------------------------------------------------------

    if st.button(
        "🔮 Predict Network QoS",
        use_container_width=True,
        type="primary",
        key="predict_qos_button"
    ):

        with st.spinner("Analyzing network performance..."):

            input_data = pd.DataFrame([{
                "rsrp": rsrp,
                "rsrq": rsrq,
                "rssi": rssi,
                "sinr": sinr,
                "delay": delay,
                "throughput_uplink": throughput_uplink,
                "ran": ran,
                "band": band
            }])

            encoded_input = preprocessor.transform(input_data)

            scaled_input = scaler.transform(encoded_input)

            prediction = qos_model.predict(scaled_input)[0]

        # ----------------------------------------------------
        # RESULT
        # ----------------------------------------------------

        st.markdown("<br>", unsafe_allow_html=True)

        if prediction == "Good":

            st.success(
                "🟢 Predicted QoS: GOOD\n\n"
                "The network conditions indicate good overall "
                "quality of service."
            )

        elif prediction == "Moderate":

            st.warning(
                "🟡 Predicted QoS: MODERATE\n\n"
                "The network conditions indicate moderate "
                "quality of service."
            )

        else:

            st.error(
                "🔴 Predicted QoS: POOR\n\n"
                "The network conditions indicate poor "
                "quality of service."
            )

        # ----------------------------------------------------
        # INPUT SUMMARY
        # ----------------------------------------------------

        st.markdown("### 📋 Input Summary")

        summary_col1, summary_col2 = st.columns(2)

        with summary_col1:

            st.write(f"**RSRP:** {rsrp} dBm")
            st.write(f"**RSRQ:** {rsrq} dB")
            st.write(f"**RSSI:** {rssi} dBm")
            st.write(f"**SINR:** {sinr} dB")

        with summary_col2:

            st.write(f"**Delay:** {delay} ms")
            st.write(
                f"**Uplink Throughput:** "
                f"{throughput_uplink} Mbps"
            )
            st.write(f"**RAN:** {ran}")
            st.write(f"**Band:** {band}")

        st.caption(
            "Prediction generated using the trained Random Forest QoS model."
        )


# ============================================================
# HISTORY
# ============================================================

elif selected_page == "🕘 History":

    st.markdown(
        '<div class="page-title">🕘 History</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'Your recently used CNW learning activities.'
        '</div>',
        unsafe_allow_html=True
    )

    if not st.session_state.history:

        render_html(
            """
            <div class="info-card"
                 style="text-align:center;">

                <div style="
                    font-size:45px;
                ">
                    🕘
                </div>

                <h2>
                    No history yet
                </h2>

                <p>
                    Your recent Study, Exam, Viva and
                    other activities will appear here.
                </p>

            </div>
            """
        )

    else:

        for item in st.session_state.history:

            render_html(
                f"""
                <div class="info-card">

                    <div style="
                        color:#2563EB;
                        font-size:15px;
                        font-weight:800;
                    ">
                        {item.get("mode", "")}
                    </div>

                    <div style="
                        color:#172033;
                        font-size:19px;
                        font-weight:650;
                        margin-top:7px;
                    ">
                        {item.get("topic", "")}
                    </div>

                </div>
                """
            )


# ============================================================
# ABOUT
# ============================================================

elif selected_page == "ℹ️ About":

    st.markdown(
        '<div class="page-title">ℹ️ About CNW AI</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'AI-powered learning assistant for Cellular and Wireless Networks.'
        '</div>',
        unsafe_allow_html=True
    )

    render_html(
        """
        <div class="info-card">

            <h2 style="
                color:#2563EB;
            ">
                📡 CNW AI Learning Assistant
            </h2>

            <p>
                CNW AI is designed to help engineering students
                learn, revise and practice Cellular and Wireless
                Networks concepts using an AI-powered interface.
            </p>

            <h3>
                Main Features
            </h3>

            <p>
                📚 Study Mode<br>
                📝 Exam Mode<br>
                🎤 Viva Mode<br>
                📊 Presentation Generator<br>
                🧠 Mind Map Generator<br>
                🎴 Flashcards<br>
                🎯 Interactive Quiz<br>
                📈 Network Performance Predictor
            </p>

            <h3>
                Technologies Used
            </h3>

            <p>
                Python • Streamlit • Google Gemini API •
                Pandas • NumPy • Scikit-learn •
                Matplotlib • PowerPoint
            </p>

        </div>
        """
    )


# ============================================================
# TEAM
# ============================================================

elif selected_page == "👥 Team":

    st.markdown(
        '<div class="page-title">👥 Team</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="page-subtitle">'
        'CNW AI Learning Assistant project team.'
        '</div>',
        unsafe_allow_html=True
    )

    c1, c2, c3 = st.columns(
        [1, 1, 1],
        gap="medium"
    )

    with c1:

        render_html(
            """
            <div class="mode-card"
                 style="text-align:center;">

                <div class="mode-icon">
                     👩‍💻
                </div>

                <div class="mode-title">
                    Team Member 1
                </div>

                <div class="mode-description">
                    Smriti Vadgule
                </div>

            </div>
            """
        )

    with c2:

        render_html(
            """
            <div class="mode-card"
                 style="text-align:center;">

                <div class="mode-icon">
                    👩‍💻
                </div>

                <div class="mode-title">
                    Team Member 2
                </div>

                <div class="mode-description">
                    Sanskruti Yadav
                </div>

            </div>
            """
        )

    
